#!/usr/bin/env python3
"""Initial Findings (blockers only) for the Federated Learning / CICIDS2017 paper,
in the required template format: # | Issue Type | Impact | Issue Details | Recommendation."""
import docx
from docx.shared import Pt, RGBColor, Inches

FINDINGS = [
    ("Research Integrity — Results reported but not actually run", "Critical",
     "Two results presented as experimental findings are hard-coded in the notebook, not measured. "
     "The noise-level ablation (Table 5, Section V-E) only genuinely ran ONE noise setting (sigma=0.0005); the sigma=0.0001 and sigma=0.001 rows come from a line the code itself labels \"# Approximate\". "
     "The convergence curve (Figure 8, Section V-F) is generated from an array the code labels \"# Theoretical convergence data\" — the training loops never record per-round accuracy. "
     "Table 5 also mislabels the sigma=0 row with the centralized NN's 98.81% instead of standard FL's 93.10%. A reviewer or integrity check could flag this.",
     "Actually run FL at each noise level and log per-round test accuracy; rebuild Table 5 and Figure 8 from real outputs. Remove any hard-coded/approximate numbers."),

    ("Methodology — 'Fair comparison' is not actually fair", "Critical",
     "The paper's headline selling point is architectural consistency (Section III-D, V-D). But in code the federated clients are class-balanced by undersampling (equal benign/attack, discarding most benign traffic), while the centralized RF and NN use the full imbalanced data. "
     "So the '5.70% distributed training cost' (Table 4) mixes two different effects: distribution/non-IID AND a different data-balancing regime. Same network shape does not equal same data handling.",
     "Use identical data handling on both sides (either balance both or neither), then re-measure the distributed-training cost. State the sampling clearly."),

    ("Core Claim — Privacy is asserted but never measured", "Critical",
     "The central contribution is privacy, yet no privacy is quantified. A small Gaussian noise (sigma=0.0005) is added once to the averaged weights, with no privacy budget (epsilon/delta), no gradient clipping/sensitivity bound, and no attack test (e.g., gradient-inversion) showing data cannot be reconstructed. "
     "This also differs from the cited DP-SGD method (Abadi 2016), which clips and adds noise per-step during training. The title/abstract/Section III-G claim more than is shown.",
     "Either add real privacy accounting (epsilon/delta with clipping) or an empirical attack-resistance test, or soften all 'privacy-preserving' claims to 'noise-perturbed'. Reconcile with the DP-SGD method cited."),

    ("Statistical Validity — Headline privacy cost is within noise", "Major",
     "The '0.41% privacy cost' is a single-run difference. Tellingly, the private model's ROC-AUC (0.9652) is HIGHER than the non-private one (0.9565), which should not happen if the noise genuinely hurts — a sign the gap is within run-to-run variance. No repeats, no error bars anywhere.",
     "Repeat each configuration with multiple seeds; report mean and standard deviation / confidence intervals. Only then interpret a 0.41% gap."),

    ("Reproducibility / Leakage — Random Forest benchmark likely inflated", "Major",
     "The 99.89% Random Forest upper bound (Section V-B.1, Table 3) is computed on data where duplicates were NOT removed (the RF code path skips de-duplication that the federated path applies). CICIDS2017 is known to contain many duplicate flows, which can place near-identical rows in both train and test and inflate accuracy. "
     "The '2.83 million records' figure matches the non-deduplicated data, contradicting Section III-C which states duplicates were removed.",
     "De-duplicate before the split for every model; recompute RF; make the record count and cleaning steps consistent between text and code."),

    ("Experimental Rigor — Single split, no cross-validation", "Major",
     "Every reported number is a single 80/20 split, single run. No cross-validation, no repeated runs, and no confidence intervals for any model or metric. This is below Q2 expectations for empirical claims.",
     "Add k-fold cross-validation or multiple seeded runs and report dispersion for all headline metrics."),

    ("Transparency — Clients silently dropped; seeds/noise unjustified", "Major",
     "The data is split by file into '8 clients' (Table 2), but the code skips any client with only one class (e.g., the all-benign Monday file), so fewer than 8 clients actually contribute — this is never stated. "
     "Random Forest and the neural networks set no random seed (limited reproducibility), and the noise scale sigma=0.0005 is chosen without justification.",
     "Report the true number of contributing clients and the skip rule; fix random seeds; justify the chosen noise scale (link it to the ablation)."),

    ("Positioning — Literature comparison does not support the method", "Major",
     "In the state-of-the-art table (Table 6, Section V-G) the proposed method's 92.70% is the lowest, and every other entry uses a different dataset, so the table cannot show competitiveness. Presented as-is it weakens rather than supports the contribution.",
     "Reframe the table explicitly as context (not a competitive ranking), or add a same-dataset baseline so a like-for-like comparison is possible."),
]

doc = docx.Document()
doc.add_paragraph("Initial Findings", style="Title")
doc.add_paragraph(
    "Federated Learning for Privacy-Preserving Big Data Analytics: "
    "A Noise-Based Privacy Framework (CICIDS2017)"
)
note = doc.add_paragraph()
r = note.add_run(
    "Pre-submission due-diligence review against Scopus Q2 expectations. "
    "Publication blockers only (Critical and Major). Findings are cross-checked against "
    "the manuscript, response sheet, and the analysis notebook."
)
r.italic = True
r.font.size = Pt(10)
doc.add_paragraph("")

table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
for i, h in enumerate(["#", "Issue Type", "Impact", "Issue Details", "Recommendation"]):
    p = table.rows[0].cells[i].paragraphs[0]
    run = p.add_run(h); run.bold = True; run.font.size = Pt(10)

impact_color = {"Critical": RGBColor(0xC0, 0x00, 0x00), "Major": RGBColor(0xB8, 0x6A, 0x00)}
for idx, (itype, impact, details, rec) in enumerate(FINDINGS, start=1):
    c = table.add_row().cells
    c[0].text = str(idx); c[1].text = itype
    pr = c[2].paragraphs[0].add_run(impact); pr.bold = True
    pr.font.color.rgb = impact_color.get(impact, RGBColor(0, 0, 0))
    c[3].text = details; c[4].text = rec
    for ci in (0, 1, 3, 4):
        for p in c[ci].paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)

for row in table.rows:
    for i, w in enumerate([Inches(0.3), Inches(1.7), Inches(0.7), Inches(3.4), Inches(2.9)]):
        row.cells[i].width = w

out = "XAI_Initial_Findings_FederatedLearning.docx"
doc.save(out)
print("wrote", out, "with", len(FINDINGS), "findings")
