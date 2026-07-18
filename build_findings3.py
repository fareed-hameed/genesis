#!/usr/bin/env python3
"""Initial Findings (blockers only) for the Energy-Efficient AI cloud-database paper,
in the required template format: # | Issue Type | Impact | Issue Details | Recommendation."""
import docx
from docx.shared import Pt, RGBColor, Inches

FINDINGS = [
    ("Validity — 100% accuracy points to label leakage", "Critical",
     "The neural network reports perfect precision, recall and F1 (1.00) with zero false positives and zero false negatives on 209,715 test samples (Table 3, Figure 5). Perfect scores almost always mean the answer is hidden in the inputs. "
     "The target 'optimal_decision' is derived from the same quantities fed to the model (e.g., system cost and transmission rate, which correlate -0.67 and +0.58 with the target in Section 4.3). Even plain logistic regression reaches 98.6% (Table 4), confirming the label is almost fully determined by the inputs.",
     "State exactly how 'optimal_decision' is generated, and remove any feature that is derived from, or used to compute, the label. Re-run and report realistic accuracy. If the label is a fixed rule, a classifier is not needed."),

    ("Scope — Synthetic data and synthetic target; no real optimization shown", "Critical",
     "The dataset and the 'optimal' label are both simulation outputs (Section 3.2), and the energy figures come from the same simulation formulas. So the study shows a network can reproduce a rule it was given, not that it saves real energy in a real system.",
     "Validate on real IoT/cloud measurements or a testbed, or clearly reframe the paper as a simulation study and soften every real-world energy claim."),

    ("Core Claim — Energy-saving result collapses and contradicts itself", "Critical",
     "The headline energy saving is 0.066% (Table 5). The trade-off table (Table 6) then shows the neural network actually uses MORE energy (196.90 J vs 196.72 J) and MORE latency (51.99 ms vs 51.77 ms) than the baselines. The two tables disagree, and the central 'energy efficiency' contribution is effectively zero.",
     "Reconcile Tables 5 and 6. If there is no real saving, present the contribution honestly rather than describing 0.066% as significant."),

    ("Internal Consistency — Record count does not match (1.5M vs 1.05M)", "Major",
     "The abstract, Section 1 and Section 3.2 state about 1.5 million records, but Table 2 shows exactly 1,048,575 for every field. That figure is the classic spreadsheet row limit (about 1.05 million), which suggests the file was opened in Excel and truncated.",
     "Use the full dataset and make the record count consistent across abstract, text and tables; confirm no rows were silently dropped."),

    ("Reproducibility — Label rule undefined and no code shared", "Major",
     "The entire result depends on how 'optimal_decision' is created, yet this rule is never given and no notebook or code accompanies the paper. A 100% result cannot be judged or reproduced without it.",
     "Provide the exact label-generation formula plus the code and data (or a clear data-availability statement)."),

    ("Reproducibility — Model and experiment details missing", "Major",
     "There is no neural network architecture (layers, neurons, epochs, learning rate), no train/validation/test split sizes, and no random seed. The energy and cost equations in Section 3 are generic and are never actually connected to what the model computes or measures.",
     "Add a full model and configuration table, state the split sizes and seed, and tie the energy equations to the values actually reported."),

    ("Interpretation — 'Neural network beats baselines' is not real evidence", "Major",
     "On a leaked target, 100% versus 98.6% (Table 4) is a meaningless gap, so the claim that the neural network captures complex nonlinear interactions is unsupported.",
     "After fixing the leakage, re-compare all models on an honest task and keep the claims proportional to the true gap."),

    ("Results — Scalability numbers are inconsistent", "Major",
     "The inference-time table (Table 7) is non-monotonic: 0.163 s for 1,000 samples but 0.106 s for 5,000, then rising again. This is not the clean near-linear scaling the text claims.",
     "Re-measure inference time with warm-up excluded and averaged over repeats; report the method so the scaling claim is credible."),
]

doc = docx.Document()
doc.add_paragraph("Initial Findings", style="Title")
doc.add_paragraph(
    "Energy-Efficient AI-Optimized Database Operations in "
    "Sustainable Cloud Computing Environments"
)
note = doc.add_paragraph()
r = note.add_run(
    "Pre-submission due-diligence review against Scopus Q2 expectations. "
    "Publication blockers only (Critical and Major), cross-checked against the manuscript "
    "tables and response sheet."
)
r.italic = True
r.font.size = Pt(10)
doc.add_paragraph("")

table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
for i, h in enumerate(["#", "Issue Type", "Impact", "Issue Details", "Recommendation"]):
    p = table.rows[0].cells[i].paragraphs[0]
    run = p.add_run(h); run.bold = True; run.font.size = Pt(10)

impact_color = {"Critical": RGBColor(0xC0, 0x00, 0x00), "Major": RGBColor(0xB8, 0x6A, 0x00)}
for idx, (itype, impact, details, rec) in enumerate(FINDINGS, start=1):
    c = table.add_row().cells
    c[0].text = str(idx); c[1].text = itype
    pr = c[2].paragraphs[0].add_run(impact); pr.bold = True
    pr.font.color.rgb = impact_color.get(impact, RGBColor(0, 0, 0))
    c[3].text = details; c[4].text = rec
    for ci in (0, 1, 3, 4):
        for p in c[ci].paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)

for row in table.rows:
    for i, w in enumerate([Inches(0.3), Inches(1.7), Inches(0.7), Inches(3.4), Inches(2.9)]):
        row.cells[i].width = w

out = "XAI_Initial_Findings_EnergyEfficientAI.docx"
doc.save(out)
print("wrote", out, "with", len(FINDINGS), "findings")
