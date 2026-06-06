#!/usr/bin/env python3
"""Generate the XAI Initial Findings document in the required template format
(Title + subtitle + 5-column issues table: # | Issue Type | Impact | Issue Details | Recommendation).
"""
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# (Issue Type, Impact, Issue Details, Recommendation)
FINDINGS = [
    ("Experimental Design — Confounded core comparison", "Critical",
     "The headline CECI finding (Abstract Feature Regime ρ = 0.61–0.72 vs. Semantic Feature Regime ρ = 0.83–0.89) rests on exactly ONE dataset per regime. \"Abstract vs. semantic feature nature\" is therefore fully confounded with dataset size (284,807 vs. 303 instances), feature count (30 vs. 13), class imbalance (577:1 vs. 1.1:1), application domain, and the number of LIME perturbations. The causal claim that feature semantics drive explanation consistency cannot be isolated from these co-varying factors, and n = 1 per regime supports no generalisation.",
     "Evaluate at least 3–5 datasets per regime spanning varied sizes, dimensionality, and balance. Use a mixed-effects / regression design that isolates semantic density (FSD) from sample size and dimensionality, or hold those factors constant across regimes."),

    ("Experimental Design — Unvalidated domain classifier", "Critical",
     "DAEA is described as a module that \"automatically learns/classifies\" the domain regime, yet it is only exercised on two datasets sitting at opposite structural extremes (FSD 0.07 vs. 0.92; DPS 0.93 vs. 0.00). There is no held-out evaluation of the classifier, no borderline/mixed cases, and no sensitivity analysis of the decision thresholds, which appear hand-tuned to separate these two points. A binary decision demonstrated on two separable points is not evidence of a working classifier.",
     "Assemble a labelled corpus of many datasets, report classifier accuracy/ROC and confusion on held-out datasets, include mixed/ambiguous cases, and provide a threshold sensitivity analysis."),

    ("Experimental Design — No ablation of the central contribution", "Critical",
     "DAEA's regime-adaptive explanation generation is the paper's primary architectural novelty, but it is never compared against a non-adaptive baseline on any quantitative outcome. CECI is computed identically with or without DAEA (it depends only on SHAP/LIME, not on the regime adaptations), so no reported number actually demonstrates that the adaptive behaviour improves anything. Claimed gains in \"usability\" and \"practical utility\" are asserted, not measured.",
     "Add a controlled ablation (standard SHAP/LIME vs. DAEA-adapted) and report a measurable explanation-quality outcome (e.g., fidelity, sparsity, stability, or a human-rated utility study) that DAEA is meant to improve."),

    ("Reproducibility — DAEA scoring is underspecified", "Critical",
     "The Domain Regime Score (DRS) is said to aggregate four discriminants (FSD, DPS, CIR, FICE), but no formula, weighting scheme, or decision threshold is given. FSD relies on \"ontological feature matching against domain knowledge bases\" with no named ontology/knowledge base or matching procedure; FICE and DPS computations are not defined. The core mechanism is not reproducible.",
     "Provide exact definitions/formulas for FSD, DPS, CIR, FICE, the DRS aggregation (weights) and thresholds, the knowledge base/ontology used, and pseudocode or released code."),

    ("Statistical Validity — No uncertainty or significance testing", "Major",
     "CECI is reported as a single point Spearman ρ per cell with no confidence intervals and no significance test of the AFR-vs-SFR difference. Performance is described as \"mean ± SD over folds,\" yet the tables report only means (no SD). The central regime-difference claim has no inferential support.",
     "Report bootstrap CIs for every CECI value, formally test the AFR/SFR difference (with n and p), and add SD/CIs to all performance and calibration tables."),

    ("Internal Consistency — Cross-validation vs. single split", "Major",
     "Methods state five-fold stratified CV (mean ± SD), but the reported confusion matrices reconcile exactly to a single 70/30 split: XGBoost fraud counts (TP 116, FP 9, FN 32) reproduce precision 0.928 and recall 0.784 from Table 4 to three decimals, and heart counts (78 correct / 91) reproduce accuracy 0.857. This suggests the \"CV\" tables may actually report one split.",
     "State a single, consistent evaluation protocol. If CV is used, report per-fold dispersion and ensure confusion matrices are aggregated over folds (or clearly labelled as one representative split)."),

    ("Internal Consistency — Misstated ECE improvement range", "Major",
     "The abstract/conclusion claim ECE reductions of \"43.4%–65.3%.\" Recomputing every row of Table 5 gives a range of 56.6% (LR-Fraud: 0.0431→0.0187) to 65.3% (XGB-Fraud). No model yields 43.4%; the stated lower bound is unsupported by the reported data.",
     "Correct the reported range to match Table 5 (≈56.6–65.3%) and recheck all derived percentages throughout."),

    ("Internal Consistency — Conflicting fraud feature count", "Major",
     "Section 1 (Introduction) states the Credit Card Fraud dataset has \"13 features,\" while Section 3.3.1 and Table 1 state 30 (28 PCA components V1–V28 plus Amount and Time). The two figures contradict each other.",
     "Standardise to 30 features (or specify the subset used) everywhere, and verify all dataset descriptors against the data actually used."),

    ("Metric Validity — Wrong discrimination metric for extreme imbalance", "Major",
     "Under 577:1 imbalance, AUC-ROC (reported 0.92–0.98 on fraud) is well known to be optimistic because true negatives dominate. The recommended discrimination metric for rare-event detection is Area Under the Precision–Recall Curve (AUPRC / average precision), which is absent.",
     "Add AUPRC and precision–recall curves for the fraud task and foreground them over AUC-ROC in the imbalanced setting."),

    ("Calibration Methodology — Unreliable on the small clinical set", "Major",
     "For the 303-instance Heart dataset, a 15% calibration hold-out is ≈32 instances and the test partition ≈91. Fitting isotonic regression on ≈32 points overfits, and a 10-bin ECE on ≈91 test instances (≈9 per bin) is extremely noisy. The heart-domain calibration claims rest on statistically fragile estimates.",
     "Prefer Platt scaling for small samples, use fewer/adaptive bins, report ECE with bootstrap CIs, and consider nested cross-validation for calibration on the small dataset."),

    ("Novelty & Positioning — CECI overlaps prior explanation-agreement work", "Major",
     "CECI is a Spearman rank correlation between SHAP global and aggregated LIME local rankings. Rank-correlation/agreement and stability measures between explanation methods are an established line of work (e.g., the explanation-disagreement and explanation-robustness/stability literature). The \"novel metric\" claim is not positioned against, or differentiated from, this body of work.",
     "Cite and contrast prior explanation (dis)agreement/consistency and stability metrics, and articulate precisely what CECI adds beyond them."),

    ("Reproducibility — LIME aggregation and sampling undefined", "Major",
     "CECI correlates SHAP global rankings with \"aggregated LIME local rankings\" over 200 instances, but the aggregation rule (mean |weight|, sign handling, frequency, or rank averaging), the instance-selection procedure, and LIME's known run-to-run instability are not specified or quantified.",
     "Define the LIME-to-global aggregation precisely, describe instance sampling, and report LIME variance across seeds (and its propagation into CECI)."),

    ("Overclaiming / Scope — \"Big Data\" framing not exercised", "Major",
     "Title, abstract and Section 1–2 foreground big data, the 5Vs, and Hadoop/Spark, but no distributed or streaming computation is performed; the largest dataset (284,807 rows) fits comfortably in memory and distributed deployment is deferred to Future Work. The framing materially overstates the demonstrated scope.",
     "Either demonstrate scalability (distributed/streaming experiments, runtime/throughput) or substantially soften the big-data framing to match what is actually shown."),

    ("Overclaiming / Scope — \"Domain-agnostic/automatic\" contradicts ontology dependence", "Major",
     "DAEA is described as domain-agnostic and automatic, yet FSD depends on matching features to a domain ontology/knowledge base — which presupposes domain knowledge and a curated resource per domain. This is internally contradictory and unclear how it generalises to unseen domains without a pre-built ontology.",
     "Reconcile the claim: explain how semantic density is computed for an unseen domain without a hand-curated ontology, or reframe DAEA as semi-automatic/knowledge-assisted."),

    ("Experimental Design — Anecdotal clinical/financial validation", "Major",
     "The SFR contrastive-explanation evidence is a single illustrative patient, and \"clinical benchmarking\" amounts to consistency with general AHA knowledge of risk factors. There is no clinician review, no user study, and no systematic evaluation, so claims of improved clinical/operational utility are unsubstantiated.",
     "Add expert (clinician/fraud-analyst) validation or a human-subjects usability study, or explicitly downgrade the utility claims to illustrative."),

    ("Novelty & Positioning — \"First/no prior work\" rests on a thin review", "Major",
     "Strong primacy claims (\"first empirical quantification,\" \"no prior work adjusts explanations to data structure\") are supported by a five-study comparison table with no systematic-review methodology, risking missed prior art.",
     "Add a systematic review protocol (search strings, databases, inclusion/exclusion, PRISMA-style flow) to substantiate the gap and primacy claims, and broaden the comparator set."),

    ("Reproducibility — Missing artifacts and tuning protocol", "Major",
     "There is no code/data availability statement, no random seeds, no library/version reporting, and hyperparameters are stated as fixed values with no described tuning/validation procedure (e.g., no grid/Bayesian search or validation curves).",
     "Provide a public code repository, fixed seeds, software environment, and a documented hyperparameter-search protocol with validation evidence."),

    ("Calibration Methodology — SMOTE×calibration interaction unaddressed", "Minor",
     "Training on SMOTE-resampled data distorts base rates and raw probability outputs. Although the calibration set is held out and unresampled (good), the interaction between resampling, calibration, and operating thresholds on the true (imbalanced) distribution is not analysed.",
     "Discuss/validate the SMOTE×calibration interaction; compare against class-weighting (no resampling) and confirm calibration is evaluated on the natural class prior."),

    ("Scope — Binary regime oversimplification", "Minor",
     "Real datasets routinely mix semantically rich and abstract/engineered features; a hard binary AFR/SFR assignment will misclassify such mixed datasets. The paper notes continuous regime scoring only as future work, which limits the generalisability of the current claims.",
     "Acknowledge the limitation prominently and, ideally, demonstrate behaviour on at least one mixed-feature dataset."),

    ("Experimental Design — Small, dated clinical benchmark", "Minor",
     "The UCI Cleveland Heart dataset (303 instances) is small and heavily reused, limiting external validity and statistical power for the SFR conclusions.",
     "Add a larger and more contemporary clinical dataset to strengthen external validity for the semantic regime."),
]

doc = docx.Document()

# Title (matches template: 'Title' style)
doc.add_paragraph("Initial Findings", style="Title")

# Subtitle (normal paragraph, as in template)
doc.add_paragraph(
    "Domain-Adaptive Explainable AI Framework for Big Data–Driven "
    "Decision Support Systems (DAEA / XAI-DSS)"
)

# Short scope note
note = doc.add_paragraph()
r = note.add_run(
    "Pre-submission due-diligence review against Scopus Q2 expectations. "
    "Findings are ordered by severity (Critical → Major → Minor). "
    "Each item references concrete evidence from the manuscript tables/sections."
)
r.italic = True
r.font.size = Pt(10)

doc.add_paragraph("")

# Issues table
table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
headers = ["#", "Issue Type", "Impact", "Issue Details", "Recommendation"]
hdr = table.rows[0].cells
for i, h in enumerate(headers):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    run = p.add_run(h)
    run.bold = True
    run.font.size = Pt(10)

impact_color = {
    "Critical": RGBColor(0xC0, 0x00, 0x00),
    "Major": RGBColor(0xB8, 0x6A, 0x00),
    "Minor": RGBColor(0x55, 0x55, 0x55),
}

for idx, (itype, impact, details, rec) in enumerate(FINDINGS, start=1):
    cells = table.add_row().cells
    cells[0].text = str(idx)
    cells[1].text = itype
    # impact with colour
    cells[2].text = ""
    pr = cells[2].paragraphs[0].add_run(impact)
    pr.bold = True
    pr.font.color.rgb = impact_color.get(impact, RGBColor(0, 0, 0))
    cells[3].text = details
    cells[4].text = rec
    # uniform body font size
    for ci in (0, 1, 3, 4):
        for p in cells[ci].paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)

# Set column widths
from docx.shared import Inches
widths = [Inches(0.3), Inches(1.6), Inches(0.7), Inches(3.4), Inches(3.0)]
for row in table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

out = "XAI_Initial_Findings_DAEA.docx"
doc.save(out)
print("wrote", out, "with", len(FINDINGS), "findings")
